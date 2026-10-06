"""Render existing LG-03 CSVs as a readable Excel workbook and Markdown report."""
from pathlib import Path
import csv
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
METRICS = ROOT / "metrics"
LABELS = {"typed_vi": "Gõ tiếng Việt", "typed_en_reviewed": "Gõ tiếng Anh đã rà soát",
          "spoken_whisper_vi": "Nói — ASR Whisper", "spoken_phowhisper_vi": "Nói — ASR PhoWhisper"}


def read(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def style(sheet, widths, header_row=2):
    sheet.freeze_panes = f"A{header_row+1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for cell in sheet[header_row]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[header_row].height = 42
    for i, width in enumerate(widths, 1): sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet.iter_rows(min_row=header_row+1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 1: cell.fill = PatternFill("solid", fgColor="EDF3F8")
            if isinstance(cell.value, float): cell.number_format = "0.00"


def main():
    rows = read(METRICS / "LG03_model_language_input.csv")
    paired = read(METRICS / "LG03_paired_model_comparison.csv")
    wb = Workbook()
    wb.remove(wb.active)
    doc = ["# LG-03 — Bảng mô hình × ngôn ngữ × gõ/nói", "",
           "Đã đánh giá trên cùng 76 câu dev. Tập 174 câu test chưa được mã hóa hoặc đánh giá.", "",
           "M-CLIP dùng ảnh ViT-L/14 có sẵn trong M1; OpenCLIP dùng ảnh ViT-B/32 mã hóa lại.", ""]
    for source, title in [("edata", "9 camera cố định"), ("video", "Video điện thoại")]:
        subset = [r for r in rows if r["source"] == source]
        sheet = wb.create_sheet(title)
        sheet.append([f"LG-03 — {title} — DEV — 06/10/2026"])
        sheet.merge_cells("A1:N1")
        sheet["A1"].font = Font(bold=True, size=14, color="17365D")
        sheet.append(["Mô hình", "Ngôn ngữ", "Đầu vào", "ASR", "n", "Gallery track", "Recall@1 (%)", "Recall@5 (%)",
                      "Recall@10 (%)", "mAP (%)", "Text p50 (ms)", "Text p95 (ms)", "Pipeline p50 (ms)", "Pipeline p95 (ms)"])
        doc.extend([f"## {title}", "", "| Mô hình | Truy vấn | n | R@1 | R@5 | R@10 | mAP |",
                    "|---|---|---:|---:|---:|---:|---:|"])
        for r in subset:
            sheet.append([r["model"], "Tiếng Anh" if r["language"] == "en" else "Tiếng Việt",
                          "Gõ" if r["input_mode"] == "typed" else "Nói (bản chép ASR)", r["asr"], int(r["n"]), int(r["gallery_tracks"])] +
                         [float(r[k]) for k in ["Recall@1_pct", "Recall@5_pct", "Recall@10_pct", "mAP_pct", "text_p50_ms", "text_p95_ms", "pipeline_p50_ms", "pipeline_p95_ms"]])
            doc.append(f"| {r['model']} | {LABELS[r['variant']]} | {r['n']} | " + " | ".join(f"{float(r[k]):.2f}%" for k in ["Recall@1_pct", "Recall@5_pct", "Recall@10_pct", "mAP_pct"]) + " |")
        style(sheet, [15, 14, 25, 16, 8, 15] + [16]*8)
        doc.append("")
    sheet = wb.create_sheet("Khoảng tin cậy")
    sheet.append(["CI 95% theo track; giữ cố định camera dev; bootstrap 1000 lần"])
    sheet.append(["Nguồn", "Mô hình", "Truy vấn", "n", "Metric", "Giá trị (%)", "CI dưới (%)", "CI trên (%)", "Diễn giải"])
    for r in rows:
        for m in ["Recall@1", "Recall@5", "Recall@10", "mAP"]:
            text = r[m+"_interpretation"]
            message = "CI chứa chance; chưa đủ bằng chứng theo CI đã chọn" if text.startswith("CI overlaps") else ("CI trên chance" if text == "CI above chance" else "CI dưới chance")
            sheet.append([r["source"], r["model"], LABELS[r["variant"]], int(r["n"]), m,
                          float(r[m+"_pct"]), float(r[m+"_ci95_low_pct"]), float(r[m+"_ci95_high_pct"]), message])
    style(sheet, [12,15,30,8,15,15,15,15,55])
    sheet = wb.create_sheet("So sánh cặp")
    sheet.append(["Chênh lệch theo cặp trên cùng truy vấn; CI chứa 0 thì chưa đủ bằng chứng theo CI"])
    sheet.append(["So sánh (phải trừ trái)", "Nguồn", "Metric", "n", "Chênh lệch (điểm %)", "CI dưới (điểm %)", "CI trên (điểm %)", "Diễn giải"])
    for r in paired:
        sheet.append([r["contrast"],r["source"],r["metric"],int(r["n"]),float(r["difference_percentage_points"]),float(r["ci95_low_pp"]),float(r["ci95_high_pp"]),r["interpretation"]])
    style(sheet,[60,12,15,8,22,22,22,55])
    sheet = wb.create_sheet("Bản dịch dev")
    sheet.append(["Rà soát bởi Codex; chưa có human sign-off; không sửa câu ASR"])
    review = read(ROOT / "configs/LG03_translation_review.csv")
    sheet.append(["annot_id", "Tiếng Việt", "Tiếng Anh gốc", "Tiếng Anh đã rà soát", "Trạng thái", "Ghi chú"])
    for r in review: sheet.append([int(r["annot_id"]),r["text_vi"],r["original_en"],r["reviewed_en"],r["status"],r["notes"]])
    style(sheet,[12,65,65,65,24,75])
    for i in range(3,sheet.max_row+1): sheet.row_dimensions[i].height=66
    sheet = wb.create_sheet("Lưu ý")
    notes = ["174 truy vấn test chưa được đánh giá; chỉ 76 dev được xử lý.",
             "edata: 70 câu dev, 4.272 track gallery, 9 camera; video: 6 câu dev, 32.404 track, 8 video.",
             "Mỗi câu có một relevant track; mAP tương đương MRR. Không có nhãn đầy đủ cho người khác cũng khớp mô tả.",
             "Tiếng Anh được Codex rà soát theo nghĩa câu Việt, chưa có xác nhận của người trong nhóm; xem sheet Bản dịch dev.",
             "Whisper/PhoWhisper là bản chép ASR gốc từ M1. Không chạy lại ASR; latency không gồm nhận dạng giọng nói.",
             "Hai hệ dùng backbone ảnh khác nhau. OpenCLIP ảnh autocast FP16, text FP32; M-CLIP ảnh từ M1 chưa rõ precision gốc.",
             "Pipeline dùng gallery có sẵn, không gồm detector, nhúng ảnh, đọc video hay tải mô hình. Lịch GPU riêng chưa xác minh.",
             "Video dev chỉ có 6 câu; CI bootstrap có thể suy biến, không suy từ CI [0,0] rằng xác suất thật bằng 0.",
             "So sánh cặp trên dev mang tính thăm dò; chưa xử lý đa kiểm định. Không kết luận mô hình hơn chỉ từ giá trị trung bình.",
             "CSV chuẩn nhóm giữ 17 cột; workbook này là bản trình bày. Chi tiết protocol: reports/LG03_protocol.md."]
    sheet.append(["Lưu ý khi đọc bảng LG-03"])
    sheet.append(["Nội dung"])
    for note in notes: sheet.append([note])
    style(sheet,[120])
    for i in range(3,sheet.max_row+1): sheet.row_dimensions[i].height=45
    dest = METRICS / "LG03_model_language_input.xlsx"
    wb.save(dest)
    doc.extend(["## Cách diễn giải", "", "- CI, latency và chênh lệch cặp đầy đủ nằm trong workbook và CSV.",
                "- CI metric chứa chance: chưa đủ bằng chứng theo CI đã chọn; CI chênh lệch chứa 0: chưa đủ bằng chứng phân biệt hai cấu hình theo CI đó.",
                "- Bản dịch được Codex rà soát; cần người trong nhóm xác nhận nếu yêu cầu human-reviewed.",
                "- Nói ở đây là bản chép ASR tiếng Việt, chưa có ô nói tiếng Anh. Latency không gồm chạy ASR.",
                "- Hai hệ khác backbone ảnh và precision ảnh; đây là so hệ retrieval, không cô lập ảnh hưởng của text encoder.",
                "- Xem [protocol LG-03](LG03_protocol.md) và [bản từng chỉ số xuống dòng](LG03_comparison_readable.txt).", ""])
    (ROOT / "reports/LG03_results.md").write_text("\n".join(doc),encoding="utf-8")
    print(f"Created {dest}; {len(rows)} comparison rows; {len(paired)} paired-metric rows")


if __name__ == "__main__": main()
